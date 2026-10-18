"""Build the unpaid GigMath Pro Kit workbook. Not linked from the site."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties
from openpyxl.chart import Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path

OUT = Path("downloads/pro/gigmath-pro-kit.xlsx")
NAVY = "0F2A43"
GREEN = "14A46B"
PALE = "F6F8FB"
AMBER = "FFF8E6"
LINE = "E2E8F0"
thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
header_fill = PatternFill("solid", fgColor=NAVY)
header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
title_font = Font(color=NAVY, bold=True, name="Calibri", size=16)
label_font = Font(name="Calibri", size=11)
input_fill = PatternFill("solid", fgColor="E8F8F1")
note_fill = PatternFill("solid", fgColor=AMBER)
money = '"$"#,##0.00'
miles_fmt = '#,##0.00'
pct = '0.00%'
date_fmt = 'YYYY-MM-DD'

def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def widths(ws, pairs):
    for col, width in pairs.items():
        ws.column_dimensions[col].width = width

def note(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(italic=True, color="4A5568", name="Calibri", size=10)
    ws[cell].alignment = Alignment(wrap_text=True, vertical="center")

def build():
    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True

    readme = wb.active
    readme.title = "Read Me"
    readme["A1"] = "GigMath Pro Kit"
    readme["A1"].font = title_font
    readme.merge_cells("A1:B1")
    lines = [
        "Educational spreadsheet for a US gig driver or 1099 freelancer. Not tax, legal, or financial advice.",
        "IRS business standard mileage rate for 2026, from irs.gov/tax-professionals/standard-mileage-rates:",
        "72.5 cents per mile for January 1 through June 30, 2026 (Notice 2026-10).",
        "76 cents per mile for July 1 through December 31, 2026 (Announcement 2026-11).",
        "Change the two rate cells on Settings if the IRS revises them. Do not treat those rates as cash the app paid you.",
        "Standard mileage replaces gas, oil, insurance, registration, repairs, lease payments, and depreciation. Do not also deduct those.",
        "Business parking, tolls, and the business percentage of car-loan interest can be separate. See IRS Publication 463.",
        "Commuting is not a business mile. Publication 463 treats the trip from home to the first stop, and from the last stop home, as commuting unless a listed exception applies.",
        "Self-employment tax follows Schedule SE: profit × 92.35%. If that result is under $400, the sheet shows no self-employment tax. Social Security is 12.4% up to the 2026 $184,500 wage base. Medicare is 2.9% with no cap.",
        "Half of self-employment tax is subtracted before the income-tax estimate. Additional Medicare Tax (Form 8959) and the qualified business income deduction are not calculated.",
        "The income-tax rate cell is yours. GigMath does not supply an official rate.",
        "Weekly Earnings tracks cash and a mileage-deduction estimate. Do not subtract the mileage rate from cash left.",
        "If you log fuel on Weekly Earnings, do not log the same fuel again on Expense Tracker.",
        "Quarterly Tax uses the standard-mileage profit (pay + tips − fees − mileage deduction − other deductible expenses) or the actual-cost profit, based on the method you pick. Equal quarterly amounts are a simplification, not Publication 505 annualization.",
        "2026 Form 1040-ES dates: April 15, 2026; June 15, 2026; September 15, 2026; January 15, 2027.",
        "This workbook is the Pro Kit product. It is not a public download.",
    ]
    for i, text in enumerate(lines, start=3):
        note(readme, f"A{i}", text)
        readme.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        readme.row_dimensions[i].height = 32
    widths(readme, {"A": 120, "B": 20})
    readme.row_dimensions[1].height = 24
    readme.page_setup.orientation = "landscape"
    readme.page_setup.fitToPage = True
    readme.page_setup.fitToWidth = 1
    readme.page_setup.fitToHeight = 1
    readme.sheet_properties.pageSetUpPr.fitToPage = True
    readme.oddHeader.left.text = "GigMath Pro Kit"
    readme.oddFooter.left.text = "Not tax advice. Confirm figures at IRS.gov."
    readme.print_title_rows = "1:1"
    readme.page_setup.paperSize = readme.PAPERSIZE_TABLOID
    readme.sheet_view.showGridLines = False
    readme.sheet_view.zoomScale = 120

    settings = wb.create_sheet("Settings")
    settings["A1"] = "Settings you can edit"
    settings["A1"].font = title_font
    settings.merge_cells("A1:C1")
    headers = ["Setting", "Value", "Source or instruction"]
    for col, text in enumerate(headers, 1):
        settings.cell(3, col, text)
    style_header(settings, 3, 3)
    rows = [
        ("Business mileage rate, Jan 1–Jun 30, 2026", 0.725, "IRS Notice 2026-10. Edit only if the IRS changes it."),
        ("Business mileage rate, Jul 1–Dec 31, 2026", 0.76, "IRS Announcement 2026-11. Edit only if the IRS changes it."),
        ("Your estimated federal income-tax rate", None, "Type a decimal such as 0.12 or 0.22. Leave blank to estimate self-employment tax only."),
        ("2026 Social Security wage base", 184500, "Printed on 2026 Schedule SE. Edit only if the IRS figure differs."),
    ]
    for offset, (label, value, source) in enumerate(rows):
        r = 4 + offset
        settings.cell(r, 1, label).font = label_font
        cell = settings.cell(r, 2, value if value is not None else None)
        cell.fill = input_fill
        cell.border = thin
        cell.font = Font(name="Calibri", size=12, bold=True, color=GREEN)
        settings.cell(r, 3, source).font = Font(name="Calibri", size=10, color="4A5568")
    settings["B4"].number_format = '"$"#,##0.000'
    settings["B5"].number_format = '"$"#,##0.000'
    settings["B6"].number_format = pct
    settings["B7"].number_format = '"$"#,##0'
    settings["A9"] = "Green cells are inputs. Other sheets read them. Do not type over formulas on the log sheets."
    settings["A9"].font = Font(italic=True, name="Calibri", size=10)
    settings.merge_cells("A9:C9")
    widths(settings, {"A": 62, "B": 18, "C": 78})
    settings.row_dimensions[3].height = 22
    settings.auto_filter.ref = "A3:C7"
    settings.freeze_panes = "A4"
    settings.page_setup.orientation = "landscape"
    settings.page_setup.fitToPage = True
    settings.page_setup.fitToWidth = 1
    settings.page_setup.fitToHeight = 1
    settings.sheet_properties.pageSetUpPr.fitToPage = True
    settings.print_title_rows = "1:3"
    settings.oddHeader.left.text = "GigMath Pro Kit — Settings"
    settings.page_setup.paperSize = settings.PAPERSIZE_TABLOID
    settings.sheet_view.zoomScale = 120

    mile = wb.create_sheet("Mileage Log")
    mile["A1"] = "Mileage log"
    mile["A1"].font = title_font
    mile.merge_cells("A1:I1")
    note(mile, "A2", "Enter the date, odometer readings, destination, and business purpose. Miles and the deduction calculate themselves. Leave commuting trips off this sheet. The rate switches on July 1, 2026 using Settings.")
    mile.merge_cells("A2:I2")
    mile.row_dimensions[2].height = 32
    mile_headers = ["Date", "Platform", "Start odometer", "End odometer", "Miles", "Destination", "Business purpose", "IRS rate", "Deduction"]
    for col, text in enumerate(mile_headers, 1):
        mile.cell(4, col, text)
    style_header(mile, 4, 9)
    first, last = 5, 204
    for r in range(first, last + 1):
        for col in (1, 2, 3, 4, 6, 7):
            mile.cell(r, col).fill = input_fill
            mile.cell(r, col).border = thin
        mile.cell(r, 1).number_format = date_fmt
        mile.cell(r, 3).number_format = '#,##0.0'
        mile.cell(r, 4).number_format = '#,##0.0'
        mile.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})')
        mile.cell(r, 5).number_format = miles_fmt
        mile.cell(r, 8, f'=IF(A{r}="","",IF(A{r}<DATE(2026,7,1),Settings!$B$4,Settings!$B$5))')
        mile.cell(r, 8).number_format = '"$"#,##0.000'
        mile.cell(r, 9, f'=IF(OR(E{r}="",H{r}=""),"",E{r}*H{r})')
        mile.cell(r, 9).number_format = money
        for col in (5, 8, 9):
            mile.cell(r, col).border = thin
    mile["K3"] = "Total business miles"
    mile["L3"] = f"=SUM(E{first}:E{last})"
    mile["K4"] = "Total standard-mileage deduction"
    mile["L4"] = f"=SUM(I{first}:I{last})"
    mile["L3"].number_format = miles_fmt
    mile["L4"].number_format = money
    mile["K3"].font = Font(bold=True, name="Calibri")
    mile["K4"].font = Font(bold=True, name="Calibri")
    mile["L3"].font = Font(bold=True, color=GREEN, name="Calibri", size=14)
    mile["L4"].font = Font(bold=True, color=GREEN, name="Calibri", size=14)
    mile.auto_filter.ref = f"A4:I{last}"
    mile.freeze_panes = "A5"
    mile.auto_filter.ref = f"A4:I{last}"
    widths(mile, {"A": 14, "B": 16, "C": 16, "D": 16, "E": 12, "F": 28, "G": 36, "H": 12, "I": 14, "J": 3, "K": 36, "L": 16})
    mile.page_setup.orientation = "landscape"
    mile.page_setup.fitToPage = True
    mile.page_setup.fitToWidth = 1
    mile.page_setup.fitToHeight = 0
    mile.page_setup.paperSize = mile.PAPERSIZE_TABLOID
    mile.sheet_properties.pageSetUpPr.fitToPage = True
    mile.print_title_rows = "1:4"
    mile.page_setup.horizontalCentered = True
    mile.oddHeader.left.text = "GigMath Pro Kit — Mileage log"
    mile.oddFooter.right.text = "Page &P of &N"
    mile.sheet_view.zoomScale = 110
    mile.conditional_formatting.add(f"E{first}:E{last}", FormulaRule(formula=[f"E{first}<0"], fill=PatternFill("solid", fgColor="FDECEC")))

    week = wb.create_sheet("Weekly Earnings")
    week["A1"] = "Weekly earnings"
    week["A1"].font = title_font
    week.merge_cells("A1:L1")
    note(week, "A2", "One row per platform per week. Cash left subtracts fees and fuel you type. It does not subtract the mileage-rate estimate. That estimate is a tax figure, and only if you use the standard mileage method.")
    week.merge_cells("A2:L2")
    week.row_dimensions[2].height = 32
    week_headers = ["Week starting", "Platform", "Pay", "Tips", "Fees", "Business miles", "Online hours", "Fuel paid", "Gross", "Mileage deduction estimate", "Cash left", "Cash per online hour"]
    for col, text in enumerate(week_headers, 1):
        week.cell(4, col, text)
    style_header(week, 4, 12)
    w1, w2 = 5, 108
    for r in range(w1, w2 + 1):
        for col in range(1, 9):
            week.cell(r, col).fill = input_fill
            week.cell(r, col).border = thin
        week.cell(r, 1).number_format = date_fmt
        for col in (3, 4, 5, 8):
            week.cell(r, col).number_format = money
        week.cell(r, 6).number_format = miles_fmt
        week.cell(r, 7).number_format = '0.00'
        week.cell(r, 9, f'=IF(COUNT(C{r}:D{r})=0,"",C{r}+D{r})')
        week.cell(r, 10, f'=IF(OR(A{r}="",F{r}=""),"",F{r}*IF(A{r}<DATE(2026,7,1),Settings!$B$4,Settings!$B$5))')
        week.cell(r, 11, f'=IF(I{r}="","",I{r}-SUM(E{r},H{r}))')
        week.cell(r, 12, f'=IF(OR(K{r}="",G{r}="",G{r}=0),"",K{r}/G{r})')
        for col in range(9, 13):
            week.cell(r, col).border = thin
            week.cell(r, col).number_format = money
    week["A110"] = "Totals"
    week["A110"].font = Font(bold=True, color="FFFFFF", name="Calibri")
    week["A110"].fill = header_fill
    for col, letter in enumerate("CDEFHIJK", 3):
        # map carefully
        pass
    for col in (3, 4, 5, 6, 7, 8, 9, 10, 11):
        cell = week.cell(110, col, f"=SUM({get_column_letter(col)}{w1}:{get_column_letter(col)}{w2})")
        cell.font = Font(bold=True, name="Calibri")
        cell.fill = PatternFill("solid", fgColor=PALE)
        cell.number_format = miles_fmt if col == 6 else ('0.00' if col == 7 else money)
    week.cell(110, 12, f'=IF(G110=0,"",K110/G110)')
    week.cell(110, 12).number_format = money
    week.cell(110, 12).font = Font(bold=True, name="Calibri")
    week.freeze_panes = "A5"
    week.auto_filter.ref = f"A4:L{w2}"
    widths(week, {get_column_letter(i): w for i, w in enumerate([16, 16, 12, 12, 12, 16, 14, 14, 14, 26, 14, 20], 1)})
    week.page_setup.orientation = "landscape"
    week.page_setup.fitToPage = True
    week.page_setup.fitToWidth = 1
    week.page_setup.fitToHeight = 0
    week.page_setup.paperSize = week.PAPERSIZE_TABLOID
    week.sheet_properties.pageSetUpPr.fitToPage = True
    week.print_title_rows = "1:4"
    week.oddHeader.left.text = "GigMath Pro Kit — Weekly earnings"
    week.oddFooter.right.text = "Page &P of &N"
    week.page_setup.horizontalCentered = True
    week.sheet_view.zoomScale = 110

    exp = wb.create_sheet("Expense Tracker")
    exp["A1"] = "Expense tracker"
    exp["A1"].font = title_font
    exp.merge_cells("A1:F1")
    note(exp, "A2", "Parking, tolls, the business percentage of car-loan interest, and other non-mileage costs. If you use the standard mileage method, do not list gas, insurance, repairs, or registration here — those are inside the IRS rate. If you use actual costs, record fuel either here or on Weekly Earnings, not both. Business use of 100% is 1.")
    exp.merge_cells("A2:F2")
    exp.row_dimensions[2].height = 36
    exp_headers = ["Date", "Category", "Description", "Amount", "Business use (1 = 100%)", "Deductible amount"]
    for col, text in enumerate(exp_headers, 1):
        exp.cell(4, col, text)
    style_header(exp, 4, 6)
    e1, e2 = 5, 154
    categories = '"Supplies,Phone,Parking,Tolls,Car loan interest,Insurance,Repairs,Other"'
    dv = DataValidation(type="list", formula1=categories, allow_blank=True)
    dv.error = "Pick a category"
    dv.errorTitle = "Category"
    dv.add(f"B{e1}:B{e2}")
    exp.add_data_validation(dv)
    pct_dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    pct_dv.error = "Enter a decimal from 0 to 1"
    pct_dv.add(f"E{e1}:E{e2}")
    exp.add_data_validation(pct_dv)
    for r in range(e1, e2 + 1):
        for col in range(1, 6):
            exp.cell(r, col).fill = input_fill
            exp.cell(r, col).border = thin
        exp.cell(r, 1).number_format = date_fmt
        exp.cell(r, 4).number_format = money
        exp.cell(r, 5).number_format = pct
        exp.cell(r, 6, f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})')
        exp.cell(r, 6).number_format = money
        exp.cell(r, 6).border = thin
    exp["A156"] = "Total deductible"
    exp["F156"] = f"=SUM(F{e1}:F{e2})"
    exp["A156"].font = Font(bold=True, name="Calibri")
    exp["F156"].font = Font(bold=True, color=GREEN, name="Calibri", size=14)
    exp["F156"].number_format = money
    exp.freeze_panes = "A5"
    exp.auto_filter.ref = f"A4:F{e2}"
    widths(exp, {"A": 14, "B": 22, "C": 42, "D": 14, "E": 26, "F": 20})
    exp.page_setup.orientation = "landscape"
    exp.page_setup.fitToPage = True
    exp.page_setup.fitToWidth = 1
    exp.page_setup.fitToHeight = 0
    exp.page_setup.paperSize = exp.PAPERSIZE_TABLOID
    exp.sheet_properties.pageSetUpPr.fitToPage = True
    exp.print_title_rows = "1:4"
    exp.oddHeader.left.text = "GigMath Pro Kit — Expenses"
    exp.oddFooter.right.text = "Page &P of &N"
    exp.page_setup.horizontalCentered = True
    exp.sheet_view.zoomScale = 110

    tax = wb.create_sheet("Quarterly Tax")
    tax["A1"] = "Quarterly tax estimator"
    tax["A1"].font = title_font
    tax.merge_cells("A1:C1")
    note(tax, "A2", "Pick a car method. Standard subtracts the mileage-log deduction and does not subtract weekly fuel. Actual subtracts weekly fuel and does not subtract the mileage rate. Expense Tracker is subtracted either way, so do not duplicate a cost. This is a federal estimate, not a filed Form 1040-ES.")
    tax.merge_cells("A2:C2")
    tax.row_dimensions[2].height = 36
    tax["A4"] = "Car method (type standard or actual)"
    tax["B4"] = "standard"
    tax["B4"].fill = input_fill
    tax["B4"].font = Font(bold=True, name="Calibri", color=GREEN)
    tax["B4"].border = thin
    method = DataValidation(type="list", formula1='"standard,actual"', allow_blank=False)
    method.add("B4")
    tax.add_data_validation(method)
    tax["A5"] = "W-2 Social Security wages (Form W-2, boxes 3 and 7)"
    tax["B5"] = 0
    tax["B5"].fill = input_fill
    tax["B5"].number_format = money
    tax["B5"].border = thin
    labels = {
        7: ("Gross pay + tips", "='Weekly Earnings'!C110+'Weekly Earnings'!D110"),
        8: ("Fees", "='Weekly Earnings'!E110"),
        9: ("Mileage deduction (standard method)", "='Mileage Log'!L4"),
        10: ("Fuel typed on Weekly Earnings (actual method)", "='Weekly Earnings'!H110"),
        11: ("Other deductible expenses", "=ExpenseTrackerTotal"),
        12: ("Net profit used for this estimate", '=IF(B4="actual",B7-B8-B10-B11,B7-B8-B9-B11)'),
        14: ("Net earnings (profit × 92.35%)", "=B12*0.9235"),
        15: ("Self-employment tax applies", '=IF(B14<400,"No — under $400","Yes")'),
        16: ("Social Security room left under the wage base", "=MAX(0,Settings!B7-B5)"),
        17: ("Social Security portion (12.4%)", '=IF(B14<400,0,MIN(B14,B16)*0.124)'),
        18: ("Medicare portion (2.9%)", "=IF(B14<400,0,B14*0.029)"),
        19: ("Self-employment tax", "=B17+B18"),
        20: ("Deduction for half of self-employment tax", "=B19*0.5"),
        21: ("Income-tax set-aside at your Settings rate", '=IF(Settings!B6="",0,MAX(0,B12-B20)*Settings!B6)'),
        22: ("Federal set-aside estimate", "=B19+B21"),
        23: ("Equal amount for each of 4 due dates", "=B22/4"),
    }
    for row, (label, formula) in labels.items():
        tax.cell(row, 1, label).font = label_font
        tax.cell(row, 2, formula).border = thin
        if row != 15:
            tax.cell(row, 2).number_format = money
    tax["C9"] = "From the Mileage Log total. Weekly miles are only an estimate until you log the trips."
    tax["C9"].alignment = Alignment(wrap_text=True, vertical="center")
    for row in (12, 19, 22, 23):
        tax.cell(row, 2).font = Font(bold=True, color=GREEN, name="Calibri", size=13)
    tax["A25"] = "Form 1040-ES due dates for calendar year 2026"
    tax["A25"].font = Font(bold=True, name="Calibri", color=NAVY, size=13)
    for col, text in enumerate(["Payment", "Due date", "Equal estimate"], 1):
        tax.cell(26, col, text)
    style_header(tax, 26, 3)
    dues = [("1", "2026-04-15"), ("2", "2026-06-15"), ("3", "2026-09-15"), ("4", "2027-01-15")]
    for i, (n, due) in enumerate(dues):
        r = 27 + i
        tax.cell(r, 1, n).border = thin
        tax.cell(r, 2, due).border = thin
        tax.cell(r, 3, "=$B$23").number_format = money
        tax.cell(r, 3).border = thin
    tax["A32"] = "You can skip the January 15, 2027 payment if you file the 2026 return by February 1, 2027 and pay the balance due. Equal amounts ignore the annualized method."
    tax["A32"].font = Font(italic=True, name="Calibri", size=10, color="4A5568")
    tax.merge_cells("A32:C32")
    tax.row_dimensions[32].height = 32
    widths(tax, {"A": 68, "B": 28, "C": 42})
    tax.freeze_panes = "A4"
    tax.page_setup.orientation = "landscape"
    tax.page_setup.fitToPage = True
    tax.page_setup.fitToWidth = 1
    tax.page_setup.fitToHeight = 1
    tax.page_setup.paperSize = tax.PAPERSIZE_TABLOID
    tax.sheet_properties.pageSetUpPr.fitToPage = True
    tax.print_title_rows = "1:2"
    tax.oddHeader.left.text = "GigMath Pro Kit — Quarterly tax"
    tax.oddFooter.left.text = "Not a filed Form 1040-ES. Confirm at IRS.gov."
    tax.page_setup.horizontalCentered = True
    tax.sheet_view.zoomScale = 120
    tax.sheet_view.showGridLines = False

    # Defined name for the expense total so the quarterly sheet stays readable.
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names.add(DefinedName(name="ExpenseTrackerTotal", attr_text="'Expense Tracker'!$F$156"))

    for ws in wb.worksheets:
        ws.page_setup.horizontalCentered = True
        ws.sheet_properties.tabColor = GREEN
        ws.page_margins.left = 0.5
        ws.page_margins.right = 0.5
        ws.page_margins.top = 0.75
        ws.page_margins.bottom = 0.75
        ws.sheet_view.view = "normal"
        ws.print_options.horizontalCentered = True

    readme.sheet_view.view = "normal"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT, OUT.stat().st_size)

if __name__ == "__main__":
    build()
